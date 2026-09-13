"""SPEC-002 — the as-built documents the provisioned array. Each test cites the requirement it proves.

Parsers are pinned to `tests/fixtures/rack13_array/` (rack13arcus, 2026-09-13 00:07, after run 2)."""

from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn

from alletra_onboard.application.documents import asbuilt_parse as ap
from alletra_onboard.application.documents.asbuilt import AsBuiltData, generate_asbuilt
from alletra_onboard.domain.provisioning import VlunTemplate

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "rack13_array"


def _fx(name: str) -> str:
    return (FIX / name).read_text(encoding="utf-8")


def _read(path):
    doc = docx.Document(str(path))
    text = "".join(t.text or "" for t in doc.element.body.iter(qn("w:t")))
    return doc, text


def _headings(doc) -> list[tuple[str, str]]:
    return [(p.style.name, p.text.strip()) for p in doc.paragraphs
            if p.style is not None and p.style.name.lower().startswith("heading")]


def _tables_by_header(doc) -> dict[tuple[str, ...], list[list[str]]]:
    out = {}
    for t in doc.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        out.setdefault(tuple(rows[0]), []).extend(rows[1:])
    return out


def _array_data(**over) -> AsBuiltData:
    data = AsBuiltData(
        customer="ACME", serial_no="CZ2D2K014S", name="rack13arcus",
        showhost_d=_fx("showhost_d.txt"), showhostset=_fx("showhostset.txt"),
        showvv=_fx("showvv.txt"), showvvset=_fx("showvvset.txt"),
        showvlun_t=_fx("showvlun_t.txt"), showvlun_a=_fx("showvlun_a.txt"),
    )
    for k, v in over.items():
        setattr(data, k, v)
    return data


# ------------------------------------------------------------------ R2, R7, R10 — showvv

def test_parse_showvv_pins_the_capture():
    rows = ap.parse_showvv(_fx("showvv.txt"))
    assert len(rows) == 51                                       # "51 total"
    by = {r["Name"]: r for r in rows}
    v = by["zz_t2_vol02"]
    assert v["Prov"] == "tdvv" and v["Dedup"] == "Yes" and v["VSize(MiB)"] == "1024" and v["Type"] == "base"
    snap = by["vol1.0.260910000000"]
    assert snap["Type"] == "vcopy" and snap["CopyOf"] == "vol1.0" and snap["Prov"] == "snp"
    assert by[".mgmtdata"]["Prov"] == "full"
    assert by["zz_t2_vol03"]["VSize(MiB)"] == "2048"


def test_parse_showvv_showcols_is_header_driven():
    text = (
        "Id Name        Prov Type UsrCPG SnpCPG VSize_MB\n"
        "92 zz_t2_vol01 tpvv base SSD_r6 SSD_r6     1024\n"
        "95 odd_one     tpvv base NL_r6  --         2048\n"
        "--------------------------------------------------\n"
        " 2 total                                  3072\n"
    )
    rows = ap.parse_showvv(text)
    assert [r["Name"] for r in rows] == ["zz_t2_vol01", "odd_one"]
    assert rows[0]["UsrCPG"] == "SSD_r6" and rows[1]["UsrCPG"] == "NL_r6" and rows[0]["VSize_MB"] == "1024"
    assert ap.parse_showvv("") == []


# ------------------------------------------------------------------ R1, R2, R10 — sets

def test_parse_cli_sets_pins_hostset_and_vvset_captures():
    hs = ap.parse_cli_sets(_fx("showhostset.txt"))
    assert len(hs) == 6
    assert hs["zz_t2_hs"] == ["10.132.30.136"]
    assert hs["grp3_alletra"] == ["grp3_vmenode1", "grp3_vmenode2", "grp3_vmenode3"]

    vs = ap.parse_cli_sets(_fx("showvvset.txt"))
    assert len(vs) == 9
    assert vs["vol1"] == ["vol1.0", "vol1.1", "vol1.2"]
    assert len(vs["vol11788928808342783376"]) == 12
    assert vs["test.Snapset"] == []                              # `--` = empty
    assert vs["VortexTestinstanc-3IUw"] == ["root-KVx6"]
    assert ap.parse_cli_sets(_fx("showvvset_none.txt")) == {}
    assert ap.parse_cli_sets("") == {}


# ------------------------------------------------------------------ R3, R10 — templates

def test_parse_showvlun_templates_cli_pins_the_capture():
    t = ap.parse_showvlun_templates_cli(_fx("showvlun_t.txt"))
    assert t == [
        VlunTemplate(volume="zz_t2_vol03", target="10.132.30.136", lun=2),
        VlunTemplate(volume="zz_t2_vol01", target="set:zz_t2_hs", lun=0),
        VlunTemplate(volume="zz_t2_vol02", target="set:zz_t2_hs", lun=1),
    ]


# ------------------------------------------------------------------ R1 — hosts section

def test_hosts_section_lists_hosts_sets_and_unclaimed_logins(tmp_path):
    out, warnings = generate_asbuilt(_array_data(), tmp_path / "hosts.docx")
    doc, _text = _read(out)
    tables = _tables_by_header(doc)
    hosts = tables[("Host", "Persona", "Initiators", "Logged in on", "Host sets")]
    row = next(r for r in hosts if r[0] == "10.132.30.136")
    assert row[1] == "VMware"
    assert "10005CED8C5312A8" in row[2] and "10005CED8C5312A9" in row[2]
    assert "0:3:3" in row[3] and "1:3:4" in row[3]
    assert row[4] == "zz_t2_hs"
    assert not any(r[0] in ("--", "") for r in hosts)          # unclaimed logins are not "hosts"

    unclaimed = tables[("Initiator", "Logged in on", "Address")]
    assert any(r[0] == "10005CBA2CFF6BD0" for r in unclaimed)
    assert any(r[0].startswith("iqn.2024-12.com.hpe:dl385g10pr13u27") and r[2] == "10.132.30.135" for r in unclaimed)

    sets = tables[("Host set", "Members")]
    assert ["zz_t2_hs", "10.132.30.136"] in sets and len(sets) == 6
    assert not [w for w in warnings if "showhost" in w or "showhostset" in w]


# ------------------------------------------------------------------ R2 — volumes section

def test_volumes_section_hides_system_and_snapshot_volumes_but_counts_them(tmp_path):
    cpg = (
        "Id Name        Prov Type UsrCPG SnpCPG VSize_MB\n"
        "92 zz_t2_vol01 tpvv base SSD_r6 SSD_r6     1024\n"
        "93 zz_t2_vol02 tdvv base SSD_r6 SSD_r6     1024\n"
    )
    out, _ = generate_asbuilt(_array_data(showvv_cpg=cpg), tmp_path / "vols.docx")
    doc, text = _read(out)
    tables = _tables_by_header(doc)
    vols = tables[("Volume", "Provisioning", "Dedup", "Compression", "Size (GiB)", "CPG", "Snapshots", "VV set")]
    names = [r[0] for r in vols]
    assert "zz_t2_vol02" in names and "vol1.0" in names
    assert ".mgmtdata" not in names and "admin" not in names       # system volumes counted, not listed
    assert "vol1.0.260910000000" not in names                       # snapshots counted, not listed
    v2 = next(r for r in vols if r[0] == "zz_t2_vol02")
    assert v2[1] == "tdvv" and v2[2] == "Yes" and v2[4] == "1" and v2[5] == "SSD_r6"
    v3 = next(r for r in vols if r[0] == "zz_t2_vol03")
    assert v3[5] == "—"                                             # no CPG read for this one
    vol10 = next(r for r in vols if r[0] == "vol1.0")
    assert vol10[6] == "4" and vol10[7] == "vol1"                   # 4 snapshots, in VV set vol1
    assert "snapshot" in text and "system volume" in text          # the one-line count
    vsets = tables[("Volume set", "Members")]
    assert ["test.Snapset", "—"] in vsets or ["test.Snapset", ""] in vsets
    assert len(vsets) == 9


# ------------------------------------------------------------------ R3 — presentations section

def test_presentations_section_joins_templates_with_active_paths(tmp_path):
    out, _ = generate_asbuilt(_array_data(), tmp_path / "pres.docx")
    doc, _text = _read(out)
    pres = _tables_by_header(doc)[("LUN", "Volume", "Presented to", "Active paths")]
    assert len(pres) == 3
    by = {r[1]: r for r in pres}
    assert by["zz_t2_vol01"][0] == "0" and by["zz_t2_vol01"][2] == "host set zz_t2_hs"
    assert by["zz_t2_vol03"][2] == "host 10.132.30.136"
    assert by["zz_t2_vol01"][3].startswith("4 ") and "10.132.30.136" in by["zz_t2_vol01"][3]


def test_a_presentation_nobody_can_see_says_so(tmp_path):
    out, _ = generate_asbuilt(_array_data(showvlun_a="Lun VVName HostName -Host_WWN- Port Type Status ID\n0 total\n"),
                              tmp_path / "dead.docx")
    doc, _ = _read(out)
    pres = _tables_by_header(doc)[("LUN", "Volume", "Presented to", "Active paths")]
    assert all("none" in r[3] for r in pres)


# ------------------------------------------------------------------ R4 — zoning section

def _zoning_plan() -> dict:
    h, a = "10005CED8C5312A9", "20340002AC02F629"
    return {
        "fabrics": [{
            "fabric": "F1", "switch_host": "10.132.30.111", "active_cfg": "mycfg",
            "switch_name": "SAN6700R13U38", "fabric_name": "Training_Lab_Rack12", "switch_count": 2,
            "hosts": [{"wwpn": h, "display": "10:00:5c:ed:8c:53:12:a9", "role": "host", "fabric": "F1",
                       "host_name": "10.132.30.136", "suggested_alias": "esx136_hba1"}],
            "array_ports": [{"wwpn": a, "display": "20:34:00:02:ac:02:f6:29", "role": "array", "fabric": "F1",
                             "nsp": "0:3:4", "node": 0, "suggested_alias": "arcus_N0S3P4"}],
            "pairs": [[h, a]], "already_zoned": [], "zone_names": {},
        }],
        "offline_hosts": [], "notes": ["both switches attach at mixed parity"], "error": None,
    }


def test_zoning_section_renders_the_plan_and_the_command_set(tmp_path):
    rendered = {"commands": {"F1": ["alicreate \"esx136_hba1\", \"10:00:5c:ed:8c:53:12:a9\"",
                                    "zonecreate \"esx136_hba1_arcus_N0S3P4\", \"esx136_hba1; arcus_N0S3P4\"",
                                    "cfgadd \"mycfg\", \"esx136_hba1_arcus_N0S3P4\"", "cfgsave", "cfgenable \"mycfg\""]},
                "selected_pairs": [["10005CED8C5312A9", "20340002AC02F629"]],
                "aliases": {"10005CED8C5312A9": "esx136_hba1", "20340002AC02F629": "arcus_N0S3P4"}}
    report = {"proper": False, "zoned_hosts": ["10.132.30.136"], "unverified_hosts": ["localhost.localdomain"],
              "expected": [], "remediations": [], "notes": []}
    out, _ = generate_asbuilt(_array_data(zoning_plan=_zoning_plan(), zoning_rendered=rendered, zoning_report=report),
                              tmp_path / "zoning.docx")
    doc, text = _read(out)
    assert ("Heading 1", "SAN zoning designed in this run") in _headings(doc)
    assert "SAN6700R13U38" in text and "Training_Lab_Rack12" in text and "mycfg" in text
    zones = _tables_by_header(doc)[("Host", "HBA WWPN", "Alias", "Array port", "Array WWPN", "Alias", "Zone(s)", "Status")]
    assert zones[0][0] == "10.132.30.136" and zones[0][3] == "0:3:4" and zones[0][2] == "esx136_hba1"
    assert zones[0][7] == "command set delivered"
    assert 'zonecreate "esx136_hba1_arcus_N0S3P4"' in text
    assert "Delivered to the SAN team" in text
    assert "localhost.localdomain" in text                       # the check result closes the section


def test_zoning_section_states_when_the_step_did_not_run(tmp_path):
    out, _ = generate_asbuilt(_array_data(), tmp_path / "nozoning.docx")
    _, text = _read(out)
    assert "This run did not include the SAN zoning step." in text


# ------------------------------------------------------------------ R5 — provisioning record

def test_provisioning_record_renders_outcomes_timestamp_and_paths(tmp_path):
    result = {"outcomes": [
        {"kind": "host", "name": "10.132.30.136", "status": "updated", "detail": "adds 1 WWN"},
        {"kind": "volume", "name": "zz_t2_vol01", "status": "created", "detail": ""},
        {"kind": "vlun", "name": "zz_t2_vvs", "status": "exists", "detail": "LUN 0, LUN 1 → set:zz_t2_hs"},
    ], "error": None}
    paths = {"hosts": [{"host": "10.132.30.136", "verdict": "live", "hbas_with_paths": 2,
                        "fabrics": ["even", "odd"], "live_volumes": ["zz_t2_vol01"], "dead_volumes": [],
                        "detail": "2 HBA(s) live on both fabrics (even, odd)"}], "error": None}
    out, _ = generate_asbuilt(
        _array_data(provisioning_result=result, provisioning_applied_at="2026-09-13T00:05:12+00:00",
                    path_verification=paths),
        tmp_path / "prov.docx")
    doc, text = _read(out)
    assert ("Heading 1", "Provisioning performed in this run") in _headings(doc)
    assert "2026-09-13" in text
    rec = _tables_by_header(doc)[("Kind", "Name", "Result", "Detail")]
    assert ["host", "10.132.30.136", "Updated", "adds 1 WWN"] in rec
    assert ["vlun", "zz_t2_vvs", "Already existed", "LUN 0, LUN 1 → set:zz_t2_hs"] in rec
    verify = _tables_by_header(doc)[("Host", "Verdict", "Detail")]
    assert verify[0][0] == "10.132.30.136" and verify[0][1] == "Live"


def test_provisioning_record_states_plan_without_apply(tmp_path):
    out, _ = generate_asbuilt(_array_data(provisioning_plan={"actions": [], "notes": [], "blockers": [], "error": None}),
                              tmp_path / "planonly.docx")
    _, text = _read(out)
    assert "A plan was built but not applied." in text
    out2, _ = generate_asbuilt(_array_data(), tmp_path / "noprov.docx")
    _, text2 = _read(out2)
    assert "This run did not include the provisioning step." in text2


# ------------------------------------------------------------------ R6 — empty vs failed

def test_a_failed_read_is_named_in_the_section_and_in_the_warnings(tmp_path):
    data = _array_data(showhostset="", read_errors={"showhostset": "command timed out after 60s"})
    out, warnings = generate_asbuilt(data, tmp_path / "failed.docx")
    _, text = _read(out)
    assert "showhostset" in text and "command timed out" in text
    assert any("showhostset" in w and "timed out" in w for w in warnings)
    assert "The array reports no host sets." not in text


def test_an_empty_array_section_says_so_without_a_warning(tmp_path):
    data = _array_data(showhostset="No host set listed\n", showvvset="No vv set listed\n",
                       showvlun_t="Lun VVName HostName -Host_WWN- Port Type\n0 total\n")
    out, warnings = generate_asbuilt(data, tmp_path / "empty.docx")
    _, text = _read(out)
    assert "The array reports no host sets." in text
    assert "The array reports no volume sets." in text
    assert "The array reports no presentations." in text
    assert not [w for w in warnings if "showhostset" in w or "showvvset" in w or "showvlun" in w]


# ------------------------------------------------------------------ R9 — headings + TOC

def test_new_sections_are_heading_1_after_checkhealth_and_toc_updates_on_open(tmp_path):
    out, _ = generate_asbuilt(_array_data(), tmp_path / "toc.docx")
    doc, _ = _read(out)
    h1 = [t for s, t in _headings(doc) if s == "Heading 1"]
    i = next(k for k, t in enumerate(h1) if "checkhealth" in t.lower())
    assert h1[i + 1:] == [
        "Hosts and host sets", "Volumes and volume sets", "Presentations",
        "SAN zoning designed in this run", "Provisioning performed in this run",
    ]
    assert doc.settings.element.find(qn("w:updateFields")) is not None


# ------------------------------------------------------------------ R6, R4, R5 — the step

class _Cli:
    def __init__(self, fail: str | None = None):
        self.fail, self.ran = fail, []

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def run(self, cmd, timeout=None):
        self.ran.append(cmd)
        if self.fail and cmd.startswith(self.fail):
            raise RuntimeError(f"{cmd} failed on purpose")
        return f"# {cmd}\n"


def test_collect_continues_past_a_failing_command(monkeypatch):
    from alletra_onboard.application.documents import steps as st

    cli = _Cli(fail="showvvset")
    monkeypatch.setattr(st, "make_array_cli", lambda creds: cli)
    data = st.DocumentSteps(coord=None)._collect_asbuilt("10.0.0.1", "u", "p")
    assert "showvvset" in data.read_errors and "failed on purpose" in data.read_errors["showvvset"]
    assert data.showvlun_t.startswith("# showvlun -t")          # the reads after the failure still ran
    assert data.showvv_cpg.startswith("# showvv -showcols") and "UsrCPG" in data.showvv_cpg
    assert "showhost -d" in cli.ran and "showhostset" in cli.ran and "showvlun -a" in cli.ran


def test_run_records_take_the_latest_event_of_each_type():
    from types import SimpleNamespace

    from alletra_onboard.application.documents import steps as st

    def ev(t, data, ts):
        return SimpleNamespace(event_type=t, data=data, created_at=ts)

    events = [
        ev("zoning.plan", {"plan": {"fabrics": [], "v": 1}}, "t1"),
        ev("zoning.plan", {"plan": {"fabrics": [], "v": 2}}, "t2"),
        ev("zoning.previewed", {"report": {"proper": False}}, "t3"),
        ev("zoning.proper", {"report": {"proper": True}}, "t4"),
        ev("zoning.rendered", {"commands": {"F1": ["cfgsave"]}}, "t5"),
        ev("storage.previewed", {"plan": {"actions": []}}, "t6"),
        ev("storage.applied", {"result": {"outcomes": []}}, "t7"),
        ev("storage.paths.verified", {"verification": {"hosts": []}}, "t8"),
    ]
    coord = SimpleNamespace(list_events=lambda run_id: events)
    data = AsBuiltData()
    st.DocumentSteps(coord=coord)._run_records("r1", data)
    assert data.zoning_plan == {"fabrics": [], "v": 2}
    assert data.zoning_report == {"proper": True}               # the later check wins, whichever type
    assert data.zoning_rendered == {"commands": {"F1": ["cfgsave"]}}
    assert data.provisioning_plan == {"actions": []}
    assert data.provisioning_result == {"outcomes": []} and data.provisioning_applied_at == "t7"
    assert data.path_verification == {"hosts": []}


# ------------------------------------------------------------------ R8 — allowlist

def test_showhostset_and_showvvset_are_allowlisted_and_nothing_else_changed():
    from alletra_onboard.adapters.array.cli_client import ALLOWED_COMMANDS

    assert "showhostset" in ALLOWED_COMMANDS and "showvvset" in ALLOWED_COMMANDS
    assert not any(c.startswith(("create", "remove", "set")) for c in ALLOWED_COMMANDS)
