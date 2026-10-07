"""SPEC-015 — the Replication tab: parse, validate (one sentence per rule), derive names, compose,
and the run-creation gate. Back-compat: a workbook without the tab parses exactly as before."""

import io

import pytest
from openpyxl import load_workbook

from alletra_onboard.application.platform.init_sheet import (
    PROTECTION_COLUMNS,
    REPLICATION_SECTIONS,
    REPLICATION_SHEET_NAME,
    build_template_bytes,
    compose_workbook_bytes,
    parse_workbook_bytes,
)
from alletra_onboard.domain.models import RunMode
from alletra_onboard.domain.replication import (
    GROUP_NAME_MAX,
    TEST_GROUP,
    ProtectionRequest,
    derive_group_name,
)

_INIT = {"serial_number": "CZ2D320BT1", "mgmt_ipv4": "10.64.122.99"}
_TARGETS = {
    "prov_array_host": "10.64.122.99", "prov_array_user": "3paradm", "prov_array_password": "pw",
    "prov_vcenter_host": "vc.example.net", "prov_vcenter_user": "administrator@vsphere.local", "prov_vcenter_password": "vpw",
}
_VOLUMES = [
    {"name": "zz_rc_vol01", "size_gib": "10", "vvset": "zz_rc_vvs"},
    {"name": "zz_rc_vol02", "size_gib": "10", "vvset": "zz_rc_vvs"},
    {"name": "lone_vol", "size_gib": "1"},
]
_HOSTSETS = [{"name": "zz_rc_hs", "members": "esx1"}]
_PEER = {"peer_host": "10.64.154.190", "peer_user": "3paradm", "peer_password": "peer=pw"}


def _sheet(*, fields: dict[str, str] | None = None, rows: list[dict[str, str]] | None = None, volumes=None) -> bytes:
    """A complete workbook (Initialisation + Provisioning tabs filled) with the Replication tab as given."""
    replication = None if fields is None and rows is None else {"fields": fields or {}, "rows": rows}
    return compose_workbook_bytes(
        base=None, init=_INIT, targets=_TARGETS, volumes=volumes or _VOLUMES, hostsets=_HOSTSETS,
        replication=replication,
    )


def _parse(data: bytes, mode: RunMode = RunMode.REPLICATE):
    return parse_workbook_bytes(data, mode=mode)


# ------------------------------------------------------------------ template + back-compat

def test_template_has_the_replication_tab_with_sections_and_the_protection_table():
    wb = load_workbook(io.BytesIO(build_template_bytes()))
    assert REPLICATION_SHEET_NAME in wb.sheetnames
    ws = wb[REPLICATION_SHEET_NAME]
    texts = {str(c.value).strip().removesuffix("*").strip() for row in ws.iter_rows() for c in row if c.value}
    for _, fields in REPLICATION_SECTIONS:
        for _, label, _, _ in fields:
            assert label in texts
    for _, label, _ in PROTECTION_COLUMNS:
        assert label in texts
    # the Initialization accelerator ships no replication (as no provisioning)
    assert REPLICATION_SHEET_NAME not in load_workbook(io.BytesIO(build_template_bytes(init_only=True))).sheetnames


def test_a_workbook_without_the_tab_parses_as_before_and_has_no_replication():
    base = compose_workbook_bytes(base=None, init=_INIT, targets=_TARGETS, volumes=_VOLUMES, hostsets=_HOSTSETS)
    wb = load_workbook(io.BytesIO(base))
    del wb[REPLICATION_SHEET_NAME]
    older = io.BytesIO()
    wb.save(older)
    parsed = parse_workbook_bytes(older.getvalue(), mode=RunMode.PROVISION_ONLY)
    assert parsed.provisioning_intent is not None
    assert parsed.provisioning_intent.replication is None


def test_a_blank_replication_tab_is_not_an_error_when_replication_is_not_selected():
    parsed = parse_workbook_bytes(_sheet(), mode=RunMode.PROVISION_ONLY)
    assert parsed.provisioning_intent.replication is None


def test_replicate_mode_refuses_a_workbook_whose_tab_is_blank():
    with pytest.raises(ValueError, match="need the 'Replication' tab"):
        _parse(_sheet())


# ------------------------------------------------------------------ R1: the tab, parsed

def test_the_tab_parses_into_an_intent_with_defaults_filled():
    parsed = _parse(_sheet(fields=_PEER, rows=[{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}]))
    rep = parsed.provisioning_intent.replication
    assert rep.peer.host == "10.64.154.190" and rep.peer.username == "3paradm"
    assert rep.peer.password.get_secret_value() == "peer=pw"       # verbatim, '=' and all
    assert rep.rtt_ms is None
    assert rep.failover_test is True and rep.failover_group_name == TEST_GROUP
    [p] = rep.protections
    assert (p.vvset, p.mode, p.rpo_minutes, p.peer_cpg) == ("zz_rc_vvs", "async", 10, "SSD_r6")
    assert p.peer_vvset_name == "zz_rc_vvs_rc"
    assert p.group_name == "zz_rc_vvs_rcg"
    assert p.period_seconds == 300                                   # RPO 10 min -> period 5 min
    assert p.auto_synchronize and p.auto_recover


def test_sync_row_and_explicit_values_are_honoured():
    rows = [{"vvset": "zz_rc_vvs", "mode": "Sync", "peer_cpg": "SSD_r6", "peer_vvset": "dr_set",
             "auto_synchronize": "no", "auto_recover": "No"}]
    rep = _parse(_sheet(fields={**_PEER, "rtt_ms": "2.5", "failover_test": "no", "failover_group": "300gb"}, rows=rows)) \
        .provisioning_intent.replication
    [p] = rep.protections
    assert p.mode == "sync" and p.rpo_minutes is None and p.period_seconds is None
    assert p.peer_vvset_name == "dr_set" and not p.auto_synchronize and not p.auto_recover
    assert rep.rtt_ms == 2.5 and rep.failover_test is False and rep.failover_group_name == "300gb"


def test_the_tab_is_parsed_whenever_the_provisioning_tabs_are_even_without_a_replication_step():
    # A PROVISION_ONLY run keeps the intent so the operator can add the steps later via Custom.
    rep = parse_workbook_bytes(_sheet(fields=_PEER, rows=[{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}]), mode=RunMode.PROVISION_ONLY) \
        .provisioning_intent.replication
    assert rep is not None and rep.peer.host == "10.64.154.190"


# ------------------------------------------------------------------ R2: one sentence per rule

@pytest.mark.parametrize("fields, rows, expected", [
    ({"peer_host": "10.64.154.190"}, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}],
     "missing required fields: Peer array admin password, Peer array admin username"),
    (_PEER, [], "add at least one volume set"),
    (_PEER, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "mode": "peer persistence"}],
     "Mode must be 'async' or 'sync'.*Peer Persistence, Active Sync, SLD and 3DC are not supported"),
    (_PEER, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "mode": "bidirectional"}], "Mode must be 'async' or 'sync'"),
    (_PEER, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "rpo_minutes": "0.5"}], "RPO must be a whole number of minutes, 1 or more"),
    (_PEER, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "rpo_minutes": "0"}], "RPO must be a whole number of minutes"),
    (_PEER, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "rpo_minutes": "ten"}], "RPO must be a whole number of minutes"),
    (_PEER, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "mode": "sync"}], "is sync but 'Measured round-trip time \\(ms\\)' is blank"),
    ({**_PEER, "rtt_ms": "12"}, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "mode": "sync"}],
     "12 ms; sync replication over RCIP needs 10 ms or less"),
    ({**_PEER, "rtt_ms": "250"}, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}],
     "250 ms; async replication over RCIP needs 200 ms or less"),
    ({**_PEER, "rtt_ms": "fast"}, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}], "must be a number"),
    (_PEER, [{"vvset": "zz_rc_vvs"}], "Peer CPG is required"),
    (_PEER, [{"vvset": "not_on_volumes_tab", "peer_cpg": "SSD_r6"}], "no volume on the Volumes tab is in VV-set 'not_on_volumes_tab'"),
    (_PEER, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}, {"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}], "appears twice"),
    (_PEER, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "auto_recover": "maybe"}], "Auto recover must be 'yes' or 'no'"),
    ({**_PEER, "failover_test": "sometimes"}, [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}], "'Run the failover test' must be 'yes' or 'no'"),
])
def test_each_validation_rule_refuses_with_its_sentence(fields, rows, expected):
    with pytest.raises(ValueError, match=expected):
        _parse(_sheet(fields=fields, rows=rows))


def test_the_peer_must_be_a_different_array():
    with pytest.raises(ValueError, match="is this run's own array"):
        _parse(_sheet(fields={**_PEER, "peer_host": "10.64.122.99"}, rows=[{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6"}]))


def test_rtt_within_limits_is_accepted_for_both_modes():
    rows = [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "mode": "sync"}]
    assert _parse(_sheet(fields={**_PEER, "rtt_ms": "10"}, rows=rows)).provisioning_intent.replication.protections[0].mode == "sync"
    rows = [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "mode": "periodic"}]   # the CLI's word for async
    assert _parse(_sheet(fields={**_PEER, "rtt_ms": "200"}, rows=rows)).provisioning_intent.replication.protections[0].mode == "async"


# ------------------------------------------------------------------ R3: derived names

def test_group_names_stay_within_the_mirror_config_limit_and_are_stable():
    assert derive_group_name("zz_rc_vvs") == "zz_rc_vvs_rcg"
    long_set = "CRV_LZ_Production_Datastores_2026"
    name = derive_group_name(long_set)
    assert len(name) <= GROUP_NAME_MAX and name.endswith("_rcg")
    assert name == derive_group_name(long_set)                       # same set, same group, every time
    assert derive_group_name(long_set + "b") != name                 # a different set never collides
    assert ProtectionRequest(vvset=long_set, peer_cpg="SSD_r6").group_name == name


# ------------------------------------------------------------------ R6: compose onto an older workbook

def test_compose_adds_the_tab_to_a_base_workbook_that_predates_it_and_replaces_rows():
    base = compose_workbook_bytes(base=None, init=_INIT, targets=_TARGETS, volumes=_VOLUMES, hostsets=_HOSTSETS)
    wb = load_workbook(io.BytesIO(base))
    del wb[REPLICATION_SHEET_NAME]
    older = io.BytesIO()
    wb.save(older)

    first = compose_workbook_bytes(base=older.getvalue(), replication={
        "fields": _PEER, "rows": [{"vvset": "zz_rc_vvs", "peer_cpg": "SSD_r6", "mode": "async", "rpo_minutes": "30"}],
    })
    second = compose_workbook_bytes(base=first, replication={"rows": [{"vvset": "zz_rc_vvs", "peer_cpg": "test"}]})
    rep = _parse(second).provisioning_intent.replication
    assert rep.peer.host == "10.64.154.190"                          # fields kept from the first compose
    assert [(p.peer_cpg, p.rpo_minutes) for p in rep.protections] == [("test", 10)]   # rows REPLACED, not appended
    assert b"peer=pw" not in older.getvalue()
