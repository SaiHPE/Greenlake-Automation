"""SPEC-018 — the as-built documents what the run replicated, read back from both arrays. Each test
cites the requirement it proves. Array states come from `tests/fixtures/rc_pair/` (D22U27 <-> E18U31):
`after_apply/` (the sync run, 2026-10-09 22:07) and `after_apply_periodic_ui/` (the UI periodic run,
2026-10-10 15:4x); the plan and result JSON from `after_apply_periodic/`."""

from __future__ import annotations

import json
from pathlib import Path

import docx
from docx.oxml.ns import qn
from pydantic import SecretStr

from alletra_onboard.application.documents.asbuilt import AsBuiltData, generate_asbuilt
from alletra_onboard.application.replication.read import read_array
from alletra_onboard.domain.shared import EndpointCreds

_FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "rc_pair"
_RACK13 = Path(__file__).resolve().parents[1] / "fixtures" / "rack13_array"
_SHOWVV_A = "Name VSize_MB\nzz_rc_vol01 1024\nzz_rc_test_v01 1024\n300gb 307200\nTest 1024\n"


class _Cli:
    """The rc_pair fixtures as an array: `showrcopy*`/`showvvset` from `capture`, the rest from the 10-07 set."""

    def __init__(self, array: str, capture: str, *, showvv: str = "", fail: bool = False) -> None:
        self.array, self.capture, self.showvv, self.fail = array, capture, showvv, fail

    def __enter__(self):
        if self.fail:
            raise RuntimeError("Login failed")
        return self

    def __exit__(self, *exc):
        return None

    def run(self, command: str, timeout=None) -> str:
        if command.startswith("showvv "):
            return self.showvv
        name = command.replace(" -", "_").replace(" ", "_") + ".txt"
        if command.startswith("showrcopy") or command == "showvvset":
            return (_FIXTURES / self.capture / self.array / name).read_text(encoding="utf-8")
        return (_FIXTURES / self.array / name).read_text(encoding="utf-8")


def _creds(host: str) -> EndpointCreds:
    return EndpointCreds(host=host, username="3paradm", password=SecretStr("pw"))


def _views(capture: str, *, peer_fails: bool = False):
    a = read_array(_creds("10.64.122.99"), array_cli_factory=lambda c: _Cli("D22U27", capture, showvv=_SHOWVV_A))
    b = read_array(_creds("10.64.154.190"), array_cli_factory=lambda c: _Cli("E18U31", capture, showvv="Name VSize_MB\n", fail=peer_fails))
    return a, b


def _periodic_json(name: str) -> dict:
    return json.loads((_FIXTURES / "after_apply_periodic" / name).read_text(encoding="utf-8"))


def _tab(mode: str = "async") -> dict:
    return {"peer_host": "10.64.154.190", "failover_test": True,
            "rows": [{"vvset": "zz_rc_vvs", "mode": mode, "rpo_minutes": 10 if mode == "async" else None, "peer_cpg": "SSD_r6"}]}


def _data(**over) -> AsBuiltData:
    def fx(n: str) -> str:
        return (_RACK13 / n).read_text(encoding="utf-8")

    data = AsBuiltData(customer="ACME", serial_no="CZ2D320BT1", name="AlletraMP_D22U27",
                       showhost_d=fx("showhost_d.txt"), showhostset=fx("showhostset.txt"), showvv=fx("showvv.txt"),
                       showvvset=fx("showvvset.txt"), showvlun_t=fx("showvlun_t.txt"), showvlun_a=fx("showvlun_a.txt"))
    for k, v in over.items():
        setattr(data, k, v)
    return data


def _read(path):
    doc = docx.Document(str(path))
    return doc, "".join(t.text or "" for t in doc.element.body.iter(qn("w:t")))


def _h1(doc) -> list[str]:
    return [p.text.strip() for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"]


def _tables(doc) -> dict[tuple[str, ...], list[list[str]]]:
    out: dict[tuple[str, ...], list[list[str]]] = {}
    for t in doc.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        out.setdefault(tuple(rows[0]), []).extend(rows[1:])
    return out


# ------------------------------------------------------------------ R4: without the tab, nothing changes

def test_a_run_without_a_replication_tab_produces_todays_document(tmp_path):
    out, _ = generate_asbuilt(_data(), tmp_path / "plain.docx")
    doc, text = _read(out)
    h1 = _h1(doc)
    assert h1[-2:] == ["SAN zoning designed in this run", "Provisioning performed in this run"]
    assert "Replication configured in this run" not in h1 and "Failover test" not in h1
    assert "replicat" not in text.lower().split("provisioning performed in this run")[-1]


# ------------------------------------------------------------------ R1 + R3: the periodic run, read back

def test_replication_section_renders_partner_groups_volumes_and_removal_blocks_from_the_read_back(tmp_path):
    a, b = _views("after_apply_periodic_ui")
    result, plan = _periodic_json("result.json"), _periodic_json("plan.json")
    out, warnings = generate_asbuilt(
        _data(replication_tab=_tab(), replication_plan=plan, replication_result=result,
              replication_applied_at="2026-10-10T10:07:45+00:00", replication_primary=a, replication_peer=b),
        tmp_path / "periodic.docx")
    doc, text = _read(out)
    assert not [w for w in warnings if "replication" in w.lower()]
    assert _h1(doc)[-2:] == ["Replication configured in this run", "Failover test"]
    tables = _tables(doc)

    partner = tables[("Array", "Serial", "OS", "Remote Copy", "RCIP ports")]
    assert partner == [
        ["This array: AlletraMP_D22U27", "CZ2D320BT1", "10.5.0", "Started, Normal", "0:4:3 10.54.122.92; 1:4:3 10.54.122.93"],
        ["Peer: AlletraMP_E18U31", "CZ2D3209YV", "10.5.0", "Started, Normal", "0:4:3 10.54.154.192; 1:4:3 10.54.154.193"],
    ]
    assert ("Partnered: AlletraMP_D22U27 → AlletraMP_E18U31 via target 'AlletraMP_E18U31' (2/2 links Up); "
            "AlletraMP_E18U31 → AlletraMP_D22U27 via target 'AlletraMP_E18U31' (2/2 links Up)") in text
    links = tables[("Array", "Target", "Policy", "Link (port → peer address)", "Status")]
    assert links == [
        ["AlletraMP_D22U27", "AlletraMP_E18U31", "mirror_config", "0:4:3 → 10.54.154.192", "Up"],
        ["AlletraMP_D22U27", "AlletraMP_E18U31", "mirror_config", "1:4:3 → 10.54.154.193", "Up"],
        ["AlletraMP_E18U31", "AlletraMP_E18U31", "mirror_config", "0:4:3 → 10.54.122.92", "Up"],
        ["AlletraMP_E18U31", "AlletraMP_E18U31", "mirror_config", "1:4:3 → 10.54.122.93", "Up"],
    ]

    groups = tables[("Group", "Mode", "RPO / period", "Policies", "Role here", "Role on peer", "Status", "Peer volume set")]
    assert groups == [
        ["zz_rc_vvs_rcg", "Periodic", "RPO 10 min (period 5m)", "auto_recover, over_per_alert, auto_synchronize", "Primary", "Secondary",
         "Started · 1/1 Synced · last sync 2026-10-10 15:42:48 IST", "zz_rc_vvs_rc (1 volume(s))"],
        ["zz_rc_test_rcg", "Periodic", "RPO 10 min (period 5m)", "auto_recover, over_per_alert, auto_synchronize", "Primary", "Secondary",
         "Started · 1/1 Synced · last sync 2026-10-10 15:42:48 IST", "zz_rc_test_rc (1 volume(s))"],
    ]
    volumes = tables[("Group", "Volume here", "Size (GiB)", "Volume on peer", "Sync status", "Last sync")]
    assert volumes == [
        ["zz_rc_vvs_rcg", "zz_rc_vol01", "1", "zz_rc_vol01", "Synced", "2026-10-10 15:42:48 IST"],
        ["zz_rc_test_rcg", "zz_rc_test_v01", "1", "zz_rc_test_v01", "Synced", "2026-10-10 15:42:48 IST"],
    ]
    assert "On the peer each group's name carries the suffix .r188150" in text
    assert "<group>" not in text                                   # no placeholder-looking text for the check to flag
    assert "Configured over WSAPI at 2026-10-10 10:07 UTC" in text
    assert "No other Remote Copy groups are on AlletraMP_D22U27." in text

    # the calls, the outcomes, the two removal blocks
    i = text.index("What was run, in order")
    assert text.index("A  creatercopygroup -usr_cpg SSD_r6 AlletraMP_E18U31:SSD_r6 zz_rc_vvs_rcg AlletraMP_E18U31:periodic") > i
    outcomes = tables[("Kind", "Name", "On", "Result", "Detail")]
    assert outcomes[0] == ["Test volume", "zz_rc_test_v01", "A", "Done", "1 GiB tpvv on SSD_r6"]
    assert ["Group started", "zz_rc_test_rcg", "A", "Done", "started"] in outcomes
    j = text.index("To remove what this run created", i)
    ka, kb = text.index("# ---- A: AlletraMP_D22U27", j), text.index("# ---- B: AlletraMP_E18U31", j)
    assert ka < text.index("stoprcopygroup -f zz_rc_vvs_rcg", ka) < text.index("dismissrcopyvv -f -removevv zz_rc_vol01 zz_rc_vvs_rcg") \
        < text.index("removercopygroup -f zz_rc_vvs_rcg") < kb < text.index("removevvset -f zz_rc_vvs_rc")
    assert "may answer 'does not exist'" in text
    # R2 until SPEC-017
    assert text.index("The failover test was not run in this run.") > kb


def test_the_sync_run_lists_the_six_groups_the_run_did_not_create(tmp_path):
    a, b = _views("after_apply")
    result = {"outcomes": [{"kind": "group", "name": "zz_rc_vvs_rcg", "where": "A", "status": "created", "detail": "sync → AlletraMP_E18U31"},
                           {"kind": "group", "name": "zz_rc_test_rcg", "where": "A", "status": "created", "detail": "sync → AlletraMP_E18U31"}],
              "removals_a": ["stoprcopygroup -f zz_rc_vvs_rcg"], "removals_b": [], "notes": [], "groups_created": ["zz_rc_vvs_rcg", "zz_rc_test_rcg"], "error": None}
    plan = {"actions": [
        {"kind": "group", "name": "zz_rc_vvs_rcg", "state": "create", "detail": {"vvset": "zz_rc_vvs", "mode": "sync", "period_seconds": None, "peer_vvset": "zz_rc_vvs_rc"}, "calls": []},
        {"kind": "group", "name": "zz_rc_test_rcg", "state": "create", "detail": {"vvset": "zz_rc_test", "mode": "sync", "period_seconds": None, "peer_vvset": "zz_rc_test_rc"}, "calls": []},
    ]}
    out, _ = generate_asbuilt(_data(replication_tab=_tab("sync"), replication_plan=plan, replication_result=result,
                                    replication_primary=a, replication_peer=b), tmp_path / "sync.docx")
    doc, text = _read(out)
    tables = _tables(doc)
    groups = tables[("Group", "Mode", "RPO / period", "Policies", "Role here", "Role on peer", "Status", "Peer volume set")]
    assert groups[0][:3] == ["zz_rc_vvs_rcg", "Sync", "sync (every write)"] and groups[0][6] == "Started · 1/1 Synced"
    present = tables[("Group", "Target", "Mode", "Role here", "Status", "Volumes")]
    assert [r[0] for r in present] == ["300gb", "APP_Test", "Intern_Automation", "Intern_Automation2", "Test-RCG", "Test-RCG2"]
    assert present[0] == ["300gb", "AlletraMP_E18U31", "Sync", "Primary", "Started", "2"]
    assert "this run did not configure" in text
    assert "# ---- B:" not in text                       # no B lines, no B block


# ------------------------------------------------------------------ R3: the peer cannot be read

def test_an_unreadable_peer_is_named_and_the_document_is_still_produced(tmp_path):
    a, b = _views("after_apply_periodic_ui", peer_fails=True)
    out, warnings = generate_asbuilt(
        _data(replication_tab=_tab(), replication_plan=_periodic_json("plan.json"), replication_result=_periodic_json("result.json"),
              replication_primary=a, replication_peer=b), tmp_path / "nopeer.docx")
    doc, text = _read(out)
    assert any("could not read the peer array" in w for w in warnings)
    assert "The peer array could not be read for this section: Could not read 10.64.154.190" in text
    tables = _tables(doc)
    partner = tables[("Array", "Serial", "OS", "Remote Copy", "RCIP ports")]
    assert partner[1] == ["Peer: 10.64.154.190", "could not be read", "—", "—", "—"]
    assert "No Remote Copy partnership between the two arrays could be confirmed" in text and "one side was not readable" in text
    groups = tables[("Group", "Mode", "RPO / period", "Policies", "Role here", "Role on peer", "Status", "Peer volume set")]
    assert groups[0][5] == "could not be read" and groups[0][4] == "Primary"        # this array's side still reads
    assert groups[0][7] == "zz_rc_vvs_rc"                                             # no peer count without the peer


# ------------------------------------------------------------------ R1: planned but not applied; not run

def test_a_plan_without_an_apply_and_a_tab_without_the_step_each_say_so(tmp_path):
    a, b = _views("after_apply")                           # the six groups only, as before any apply
    a.groups = [g for g in a.groups if not g.name.startswith("zz_rc_")]
    out, _ = generate_asbuilt(_data(replication_tab=_tab(), replication_plan=_periodic_json("plan.json"),
                                    replication_primary=a, replication_peer=b), tmp_path / "planned.docx")
    _, text = _read(out)
    assert "A replication plan was built but not applied." in text
    assert "What was run, in order" not in text and "To remove what this run created" not in text
    out2, _ = generate_asbuilt(_data(replication_tab=_tab(), replication_primary=a, replication_peer=b), tmp_path / "notrun.docx")
    _, text2 = _read(out2)
    assert "The replication step was not run in this run." in text2
    assert "Partnered: AlletraMP_D22U27 → AlletraMP_E18U31" in text2            # the partnership is still documented
    assert "300gb" in text2                                                       # and so is what already replicates


# ------------------------------------------------------------------ the step: events and the read of both arrays

def test_a_run_that_found_the_groups_already_in_place_says_so(tmp_path):
    """Live 2026-10-10 19:01: the failover run re-read the arrays, every group already existed, nothing was
    applied. The section said 'built but not applied' AND 'Configured over WSAPI', and the peer set was '—'."""
    a, b = _views("after_apply")
    plan = {"actions": [
        {"kind": "group", "name": "zz_rc_vvs_rcg", "state": "exists", "reason": "Started, Primary, Sync, 1 volume(s)", "calls": [], "detail": {}},
        {"kind": "peer_vvset", "name": "zz_rc_vvs_rc", "state": "exists", "calls": [], "detail": {}},
        {"kind": "group", "name": "zz_rc_test_rcg", "state": "exists", "reason": "Started, Primary, Sync, 1 volume(s)", "calls": [], "detail": {}},
    ]}
    out, _ = generate_asbuilt(_data(replication_tab=_tab("sync"), replication_plan=plan, replication_primary=a, replication_peer=b),
                              tmp_path / "inplace.docx")
    doc, text = _read(out)
    assert "built but not applied" not in text and "Configured over WSAPI" not in text
    assert "Already in place when this run read the arrays (configured by an earlier run); this run made no replication change." in text
    groups = _tables(doc)[("Group", "Mode", "RPO / period", "Policies", "Role here", "Role on peer", "Status", "Peer volume set")]
    assert [(r[0], r[7]) for r in groups] == [("zz_rc_vvs_rcg", "zz_rc_vvs_rc (1 volume(s))"), ("zz_rc_test_rcg", "zz_rc_test_rc (1 volume(s))")]
    assert "To remove what this run created" not in text


def test_run_records_take_the_replication_plan_report_and_result():
    from types import SimpleNamespace

    from alletra_onboard.application.documents import steps as st

    def ev(t, data, ts):
        return SimpleNamespace(event_type=t, data=data, created_at=ts)

    events = [
        ev("replication.previewed", {"report": {"findings": ["x"]}, "plan": {"actions": [], "v": 1}}, "t1"),
        ev("replication.previewed", {"report": {"findings": []}, "plan": {"actions": [], "v": 2}}, "t2"),
        ev("replication.apply.failed", {"result": {"outcomes": [], "error": "boom"}}, "t3"),
        ev("replication.applied", {"result": {"outcomes": [{"kind": "group"}], "error": None}}, "t4"),
    ]
    data = AsBuiltData()
    st.DocumentSteps(coord=SimpleNamespace(list_events=lambda run_id: events))._run_records("r1", data)
    assert data.replication_plan == {"actions": [], "v": 2} and data.replication_report == {"findings": []}
    assert data.replication_result == {"outcomes": [{"kind": "group"}], "error": None}
    assert data.replication_applied_at == "t4"


def test_collect_replication_reads_both_arrays_with_the_runs_credentials():
    from alletra_onboard.application.documents import steps as st
    from alletra_onboard.domain.replication import ProtectionRequest, ReplicationIntent

    seen = []
    steps = st.DocumentSteps(coord=None, read_replication=lambda creds: seen.append(creds) or f"view:{creds.host}")
    a, b = steps._collect_replication(_creds("10.64.122.99"), _creds("10.64.154.190"))
    assert (a, b) == ("view:10.64.122.99", "view:10.64.154.190") and [c.host for c in seen] == ["10.64.122.99", "10.64.154.190"]
    intent = ReplicationIntent(peer=_creds("10.64.154.190"), rtt_ms=1.0,
                               protections=[ProtectionRequest(vvset="zz_rc_vvs", peer_cpg="SSD_r6")])
    assert st.DocumentSteps._replication_tab(intent) == {
        "peer_host": "10.64.154.190", "failover_test": True,
        "rows": [{"vvset": "zz_rc_vvs", "mode": "async", "rpo_minutes": 10, "peer_cpg": "SSD_r6"}],
    }


def test_the_failover_section_shares_the_replication_page(tmp_path):
    a, b = _views("after_apply_periodic_ui")
    out, _ = generate_asbuilt(_data(replication_tab=_tab(), replication_primary=a, replication_peer=b), tmp_path / "pages.docx")
    doc, _ = _read(out)
    h1 = {p.text.strip(): p for p in doc.paragraphs if p.style.name == "Heading 1"}
    assert h1["Replication configured in this run"].paragraph_format.page_break_before is True
    assert h1["Failover test"].paragraph_format.page_break_before is not True


# ------------------------------------------------------------------ SPEC-017 R6: the failover record

def _side(array, group, role, status, last=""):
    return {"array": array, "group": group, "present": True, "role": role, "status": status, "mode": "async", "volumes": 1, "synced": 1, "last_sync": last}


def _failover_record(result="passed", **over) -> dict:
    p, s = "AlletraMP_D22U27", "AlletraMP_E18U31"
    rec = {
        "group": "zz_rc_test_rcg", "peer_group": "zz_rc_test_rcg.r188150", "mode": "async", "primary_array": p, "peer_array": s,
        "started_at": "2026-10-10T10:20:00+00:00", "ended_at": "2026-10-10T10:26:30+00:00", "result": result, "failed_step": None,
        "time_to_failover_s": 6.2, "time_to_synced_s": 41.0,
        "data_loss_bound": "last sync before failover 2026-10-10 15:42:48 IST; failover completed at 2026-10-10T10:20:40+00:00",
        "observed_state": "", "recovery_action": "", "error": None,
        "steps": [
            {"seq": 0, "title": "Read both arrays", "where": "-", "action": "(read)", "cli": "", "started_at": "2026-10-10T10:20:00+00:00", "seconds": 5.1,
             "primary": _side(p, "zz_rc_test_rcg", "Primary", "Started", "2026-10-10 15:42:48 IST"), "peer": _side(s, "zz_rc_test_rcg.r188150", "Secondary", "Started"), "outcome": "ok"},
            {"seq": 1, "title": "Stop the group", "where": "P", "cli": "stoprcopygroup -f zz_rc_test_rcg", "started_at": "2026-10-10T10:20:06+00:00", "seconds": 3.0,
             "primary": _side(p, "zz_rc_test_rcg", "Primary", "Stopped"), "peer": _side(s, "zz_rc_test_rcg.r188150", "Secondary", "Stopped"), "outcome": "ok"},
            {"seq": 2, "title": "Fail over to the peer", "where": "S", "cli": "setrcopygroup failover -f zz_rc_test_rcg.r188150", "started_at": "2026-10-10T10:20:34+00:00", "seconds": 6.2,
             "primary": _side(p, "zz_rc_test_rcg", "Primary", "Stopped"), "peer": _side(s, "zz_rc_test_rcg.r188150", "Primary-Rev", "Stopped"), "outcome": "ok"},
            {"seq": 6, "title": "Restore the natural direction", "where": "S", "cli": "setrcopygroup restore -f zz_rc_test_rcg.r188150", "started_at": "2026-10-10T10:26:00+00:00", "seconds": 30.0,
             "primary": _side(p, "zz_rc_test_rcg", "Primary", "Started"), "peer": _side(s, "zz_rc_test_rcg.r188150", "Secondary", "Started"), "outcome": "ok"},
        ],
    }
    rec.update(over)
    return rec


def test_a_passed_failover_test_renders_the_step_table_timings_and_bound(tmp_path):
    a, b = _views("after_apply_periodic_ui")
    out, _ = generate_asbuilt(_data(replication_tab=_tab(), replication_primary=a, replication_peer=b, failover_record=_failover_record()),
                              tmp_path / "failover.docx")
    doc, text = _read(out)
    assert "The failover test was not run in this run." not in text
    assert "Result: Passed." in text
    assert "Group zz_rc_test_rcg (async) on AlletraMP_D22U27 (P), named zz_rc_test_rcg.r188150 on AlletraMP_E18U31 (S)." in text
    assert "the peer took over (Primary-Rev) 6.2 s after the failover was issued; every volume was Synced again 41 s after the recover" in text
    assert "Data-loss bound (async): last sync before failover 2026-10-10 15:42:48 IST" in text
    rows = _tables(doc)[("#", "Step", "On", "Command (CLI equivalent)", "Started", "Took", "P after", "S after", "Outcome")]
    assert rows[1] == ["1", "Stop the group", "P", "stoprcopygroup -f zz_rc_test_rcg", "2026-10-10 10:20 UTC", "3 s",
                       "Primary/Stopped · 1/1 Synced", "Secondary/Stopped · 1/1 Synced", "OK"]
    assert rows[2][7] == "Primary-Rev/Stopped · 1/1 Synced"
    assert rows[0][6] == "Primary/Started · 1/1 Synced · last sync 2026-10-10 15:42:48 IST"
    assert "ended the test in its starting state" in text


def test_a_failed_failover_test_names_the_step_the_state_and_the_way_back(tmp_path):
    a, b = _views("after_apply_periodic_ui")
    rec = _failover_record(result="failed", failed_step=4, time_to_synced_s=None,
                           error="AlletraMP_E18U31 refused the recover: HTTP 403 INV_OPERATION_RCOPY_GROUP_ROLE_CONFLICT",
                           observed_state="AlletraMP_D22U27: Primary/Stopped · 1/1 Synced · AlletraMP_E18U31: Primary-Rev/Stopped · 1/1 Synced",
                           recovery_action="On AlletraMP_E18U31: setrcopygroup recover -f zz_rc_test_rcg.r188150 ; then setrcopygroup restore -f zz_rc_test_rcg.r188150")
    rec["steps"][3]["outcome"] = "skipped"
    out, _ = generate_asbuilt(_data(replication_tab=_tab(), replication_primary=a, replication_peer=b, failover_record=rec), tmp_path / "failed.docx")
    doc, text = _read(out)
    assert "Result: Failed at step 4." in text
    assert "Stopped: AlletraMP_E18U31 refused the recover" in text
    assert "Observed when it stopped: AlletraMP_D22U27: Primary/Stopped" in text
    i = text.index("The documented way back to the normal state (not run by the tool):")
    assert text.index("On AlletraMP_E18U31: setrcopygroup recover -f zz_rc_test_rcg.r188150", i) > i
    rows = _tables(doc)[("#", "Step", "On", "Command (CLI equivalent)", "Started", "Took", "P after", "S after", "Outcome")]
    assert rows[3][8] == "Skipped" and "ended the test in its starting state" not in text


def test_run_records_take_the_failover_record_from_either_event():
    from types import SimpleNamespace

    from alletra_onboard.application.documents import steps as st

    def ev(t, data, ts):
        return SimpleNamespace(event_type=t, data=data, created_at=ts)

    events = [ev("failover.failed", {"record": {"result": "failed", "failed_step": 2}}, "t1"),
              ev("failover.completed", {"record": {"result": "passed"}}, "t2")]
    data = AsBuiltData()
    st.DocumentSteps(coord=SimpleNamespace(list_events=lambda run_id: events))._run_records("r1", data)
    assert data.failover_record == {"result": "passed"}
